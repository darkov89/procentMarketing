"""Tests for openpyxl Excel exporter."""

import openpyxl

from leadmachine.db.models import Audit, Contact, Lead
from leadmachine.exporter.excel_exporter import export_leads_to_excel


def test_export_leads_to_excel(tmp_path):
    lead1 = Lead(
        company_name="Centrum Medyczne Legnica",
        nip="6910001122",
        phone_normalized="+48768001122",
        email_primary="biuro@med-legnica.pl",
        address="ul. Złotoryjska 50",
        city="Legnica",
        distance_km=1.2,
        industry="Medycyna",
        status="qualified",
        score=75,
    )
    contact1 = Contact(
        first_name="Jan",
        last_name="Kowalski",
        role="Właściciel",
        is_primary=True,
    )
    lead1.contacts.append(contact1)
    audit1 = Audit(
        ssl_valid=True,
        is_responsive=True,
        google_rating=4.9,
        google_reviews_count=85,
    )
    lead1.audit = audit1

    out_file = tmp_path / "test_leads.xlsx"
    saved = export_leads_to_excel([lead1], out_file)

    assert saved.exists()

    # Load back with openpyxl and verify minimal required columns
    wb = openpyxl.load_workbook(saved)
    ws = wb.active

    # Row 1 headers
    headers = [cell.value for cell in ws[1]]
    assert headers[0] == "Firma"
    assert headers[1] == "Główny kontakt (właściciel)"
    assert headers[2] == "Telefon"
    assert headers[3] == "Mail firmy"

    # Row 2 values
    row2 = [cell.value for cell in ws[2]]
    assert row2[0] == "Centrum Medyczne Legnica"
    assert "Jan Kowalski" in row2[1]
    assert row2[2] == "+48768001122"
    assert row2[3] == "biuro@med-legnica.pl"

    # Verify freeze pane and auto filter
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
